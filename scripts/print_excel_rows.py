#!/usr/bin/env python3
from openpyxl import load_workbook
wb = load_workbook('/work/قائمة اسعار كايرو ميتال .xlsx', data_only=True)
sheet = wb.active
for r in range(2, min(60, sheet.max_row+1)):
    a = sheet.cell(row=r, column=1).value
    b = sheet.cell(row=r, column=2).value
    print(r, '|', repr(a), '|', repr(b))
