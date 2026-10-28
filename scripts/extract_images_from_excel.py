#!/usr/bin/env python3
import sys, os, json, re
import shutil
from pathlib import Path
try:
    from openpyxl import load_workbook
except Exception as e:
    print("ERROR: openpyxl not installed", file=sys.stderr)
    sys.exit(1)

def slugify(s: str) -> str:
    s = s.strip().lower()
    s = re.sub(r'[^a-z0-9]+', '-', s)
    s = re.sub(r'-+', '-', s)
    s = s.strip('-')
    return s or 'item'

def main():
    import argparse
    ap = argparse.ArgumentParser()
    ap.add_argument("--excel", required=True)
    ap.add_argument("--images", required=True)
    ap.add_argument("--target", required=True)
    ap.add_argument("--mapping", required=True)
    args = ap.parse_args()

    excel_path = args.excel
    images_dir = args.images
    target_dir = args.target
    mapping_path = args.mapping

    os.makedirs(target_dir, exist_ok=True)

    wb = load_workbook(excel_path, data_only=True)
    sheet = wb.active
    product_names = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        val = row[0] if row and len(row) > 0 else None
        if val is None:
            product_names.append("")
        else:
            product_names.append(str(val).strip())

    image_files = sorted([f for f in os.listdir(images_dir) if os.path.isfile(os.path.join(images_dir, f))])

    mapping = {}
    existing = set(os.listdir(target_dir))

    for idx, img in enumerate(image_files):
        if idx < len(product_names) and product_names[idx].strip():
            key = product_names[idx].strip()
        else:
            key = f"Product {idx+1}"
        slug = slugify(key)
        src = os.path.join(images_dir, img)
        ext = os.path.splitext(img)[1].lower()
        candidate = slug + ext
        unique = candidate
        n = 1
        while unique in existing:
            unique = f"{slug}-{n}{ext}"
            n += 1
        dst = os.path.join(target_dir, unique)
        shutil.copy2(src, dst)
        existing.add(unique)
        mapping.setdefault(key, []).append(unique)

    with open(mapping_path, "w", encoding="utf-8") as mf:
        json.dump(mapping, mf, ensure_ascii=False, indent=2)

    print("DONE")

if __name__ == "__main__":
    main()
