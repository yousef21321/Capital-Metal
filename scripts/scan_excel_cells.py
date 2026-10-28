#!/usr/bin/env python3
from openpyxl import load_workbook
wb = load_workbook('/work/قائمة اسعار كايرو ميتال .xlsx', data_only=True)
s = wb.active
for r in range(2, min(80, s.max_row+1)):
    row = []
    for c in range(1, 10):
        row.append(str(s.cell(row=r, column=c).value))
    print('\t'.join(row))
