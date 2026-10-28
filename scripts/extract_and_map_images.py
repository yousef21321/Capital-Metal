#!/usr/bin/env python3
import sys, os, json, re, zipfile, shutil
from pathlib import Path
try:
    from openpyxl import load_workbook
except Exception:
    print("ERROR: openpyxl not installed. Install with: python3 -m pip install --user openpyxl", file=sys.stderr)
    sys.exit(1)

def slugify(s: str) -> str:
    s = s.strip().lower()
    s = re.sub(r'[^a-z0-9]+', '-', s)
    s = re.sub(r'-+', '-', s)
    s = s.strip('-')
    return s or 'item'

def main():
    import argparse
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--target", required=True)
    ap.add_argument("--mapping", required=True)
    args = ap.parse_args()

    xlsx_path = args.xlsx
    target_dir = args.target
    mapping_path = args.mapping
    tmp_dir = "/tmp/cairo_xlsx_media"
    os.makedirs(target_dir, exist_ok=True)
    shutil.rmtree(tmp_dir, ignore_errors=True)
    os.makedirs(tmp_dir, exist_ok=True)

    # Extract images from xlsx (which is a zip)
    with zipfile.ZipFile(xlsx_path, 'r') as z:
        for name in z.namelist():
            if name.startswith("xl/media/"):
                filename = os.path.basename(name)
                if not filename:
                    continue
                with z.open(name) as src, open(os.path.join(tmp_dir, filename), 'wb') as dst:
                    dst.write(src.read())
    # Read product names from first sheet column A (excluding header)
    wb = load_workbook(xlsx_path, data_only=True)
    sheet = wb.active
    product_names = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        val = row[0] if row and len(row) > 0 else None
        product_names.append(str(val).strip()) if val is not None else product_names.append("")
    image_files = sorted([f for f in os.listdir(tmp_dir) if os.path.isfile(os.path.join(tmp_dir, f))])

    mapping = {}
    existing = set(os.listdir(target_dir))

    for idx, img in enumerate(image_files):
        key = product_names[idx].strip() if idx < len(product_names) and product_names[idx].strip() else f"Product {idx+1}"
        slug = slugify(key)
        src = os.path.join(tmp_dir, img)
        ext = os.path.splitext(img)[1].lower()
        candidate = slug + ext
        unique = candidate
        n = 1
        while unique in existing:
            unique = f"{slug}-{n}{ext}"
            n += 1
        dst = os.path.join(target_dir, unique)
        shutil.copy2(src, dst)
        existing.add(unique)
        mapping.setdefault(key, []).append(unique)

    with open(mapping_path, "w", encoding="utf-8") as mf:
        json.dump(mapping, mf, ensure_ascii=False, indent=2)

    print("DONE")

if __name__ == "__main__":
    main()
